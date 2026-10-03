#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Hermes B2B Data Scraper & ETL Pipeline
--------------------------------------
Module cào dữ liệu báo giá sản phẩm, catalog thương mại điện tử chuyên dụng cho khách hàng B2B.
Tự động trích xuất thông tin, xử lý lỗi mạng (retry/backoff), chuẩn hóa định dạng
và xuất trực tiếp ra CSV (UTF-8 BOM chống lỗi font Excel) và Excel (.xlsx format chuẩn).

Tác giả: Builder Agent - Hermes Ecosystem
Khách hàng / Doanh nghiệp: Tuấn Anh & Đối tác B2B
"""

import sys
import os
import time
import random
import logging
import argparse
from datetime import datetime
from typing import List, Dict, Any, Optional
from urllib.parse import urljoin

import requests
from bs4 import BeautifulSoup

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

# Cấu hình logging
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    datefmt="%Y-%m-%d %H:%M:%S"
)
logger = logging.getLogger("B2BScraper")

# Danh sách User-Agents thực tế để xoay vòng, giảm nguy cơ bị chặn IP
USER_AGENTS = [
    "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/128.0.0.0 Safari/537.36",
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/129.0.0.0 Safari/537.36",
    "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/127.0.0.0 Safari/537.36",
    "Mozilla/5.0 (Macintosh; Intel Mac OS X 10.15; rv:130.0) Gecko/20100101 Firefox/130.0",
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64; rv:129.0) Gecko/20100101 Firefox/129.0"
]

RATING_MAP = {
    "One": 1,
    "Two": 2,
    "Three": 3,
    "Four": 4,
    "Five": 5
}

class B2BDataScraper:
    """
    Trình thu thập dữ liệu sản phẩm & bảng giá B2B chuyên nghiệp.
    Hỗ trợ crawl đa trang, cơ chế retry tự động và xuất báo cáo đa định dạng.
    """

    def __init__(self, base_url: str = "http://books.toscrape.com/", delay_range: tuple = (0.3, 0.8)):
        self.base_url = base_url
        self.delay_range = delay_range
        self.session = requests.Session()
        self.results: List[Dict[str, Any]] = []

    def _get_headers(self) -> Dict[str, str]:
        """Tạo headers giả lập trình duyệt người dùng thật."""
        return {
            "User-Agent": random.choice(USER_AGENTS),
            "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,image/avif,image/webp,*/*;q=0.8",
            "Accept-Language": "vi-VN,vi;q=0.9,en-US;q=0.8,en;q=0.7",
            "Referer": self.base_url,
            "Connection": "keep-alive"
        }

    def fetch_page(self, url: str, max_retries: int = 3) -> Optional[str]:
        """
        Gửi yêu cầu tải mã nguồn HTML của trang với cơ chế Retry & Exponential Backoff.
        """
        for attempt in range(1, max_retries + 1):
            try:
                headers = self._get_headers()
                response = self.session.get(url, headers=headers, timeout=12)

                if response.status_code == 200:
                    response.encoding = "utf-8"
                    return response.text
                elif response.status_code in (429, 500, 502, 503, 504):
                    wait_time = (2 ** attempt) + random.uniform(0.5, 1.5)
                    logger.warning(f"Gặp mã phản hồi HTTP {response.status_code} tại {url}. Thử lại lần {attempt}/{max_retries} sau {wait_time:.1f}s...")
                    time.sleep(wait_time)
                else:
                    logger.error(f"Lỗi HTTP {response.status_code} khi tải {url}. Bỏ qua trang này.")
                    return None
            except (requests.RequestException, Exception) as e:
                wait_time = (2 ** attempt) + random.uniform(0.5, 1.5)
                logger.warning(f"Lỗi mạng ({e.__class__.__name__}: {e}) tại {url}. Thử lại lần {attempt}/{max_retries} sau {wait_time:.1f}s...")
                time.sleep(wait_time)

        logger.error(f"Thất bại hoàn toàn sau {max_retries} lần thử tải: {url}")
        return None

    def parse_product_item(self, article: BeautifulSoup, current_page_url: str) -> Optional[Dict[str, Any]]:
        """
        Bóc tách chi tiết thông tin một sản phẩm từ thẻ DOM article.
        """
        try:
            # 1. Tên sản phẩm (Title)
            h3_tag = article.find("h3")
            a_tag = h3_tag.find("a") if h3_tag else None
            if not a_tag:
                return None
            title = a_tag.get("title") or a_tag.get_text(strip=True)
            if not title or title.strip() == "" or title == "N/A":
                return None

            # 2. Đường link sản phẩm chi tiết
            rel_link = a_tag.get("href")
            if not rel_link:
                return None
            product_url = urljoin(current_page_url, rel_link)

            # 3. Giá sản phẩm (Price)
            price_elem = article.find("p", class_="price_color")
            if not price_elem:
                return None
            raw_price = price_elem.get_text(strip=True)
            numeric_price_str = "".join([c for c in raw_price if c.isdigit() or c == "."])
            try:
                price_gbp = float(numeric_price_str) if numeric_price_str else 0.0
            except ValueError:
                return None
            if price_gbp <= 0.0:
                return None

            # Giá quy đổi ước tính sang VNĐ (tỷ giá tham khảo 1 GBP = 33.000 VNĐ cho báo giá B2B)
            price_vnd = int(price_gbp * 33000)

            # 4. Trạng thái tồn kho (Stock Availability)
            instock_elem = article.find("p", class_="instock availability")
            stock_status_raw = instock_elem.get_text(strip=True) if instock_elem else "Không rõ"
            in_stock = "Còn hàng" if "In stock" in stock_status_raw else "Hết hàng"

            # 5. Đánh giá (Rating)
            rating_elem = article.find("p", class_="star-rating")
            rating_num = 0
            if rating_elem:
                classes = rating_elem.get("class", [])
                for cls in classes:
                    if cls in RATING_MAP:
                        rating_num = RATING_MAP[cls]
                        break

            # 6. Ảnh thumbnail
            img_tag = article.find("img", class_="thumbnail")
            img_rel = img_tag.get("src") if img_tag else ""
            image_url = urljoin(current_page_url, img_rel)

            import hashlib
            digest = int(hashlib.sha256(title.encode("utf-8")).hexdigest()[:8], 16)
            product_id = f"SP-{digest % 100000:05d}"

            record = {
                "Mã SP (ID)": product_id,
                "Tên sản phẩm": title,
                "Giá niêm yết (GBP)": price_gbp,
                "Giá quy đổi (VNĐ)": price_vnd,
                "Tình trạng tồn kho": in_stock,
                "Đánh giá (Sao)": rating_num,
                "Link chi tiết": product_url,
                "Link hình ảnh": image_url,
                "Thời gian thu thập": datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            }
            return record

        except Exception as e:
            logger.debug(f"Bỏ qua item lỗi định dạng: {e}")
            return None

    def scrape_pages(self, max_pages: int = 3) -> List[Dict[str, Any]]:
        """
        Crawl dữ liệu danh mục qua nhiều trang (Pagination).
        """
        logger.info(f"Bắt đầu quy trình cào dữ liệu: Tối đa {max_pages} trang từ {self.base_url}")
        current_url = self.base_url
        page_count = 0
        seen_urls = set()

        while current_url and page_count < max_pages:
            page_count += 1
            logger.info(f"Đang xử lý Trang {page_count}/{max_pages}: {current_url}")
            html_content = self.fetch_page(current_url)

            if not html_content:
                logger.warning(f"Không thể lấy nội dung Trang {page_count}. Dừng cào dữ liệu.")
                break

            soup = BeautifulSoup(html_content, "html.parser")
            articles = soup.find_all("article", class_="product_pod")

            if not articles:
                logger.info(f"Không tìm thấy sản phẩm nào trên Trang {page_count}. Đã đến trang cuối.")
                break

            for article in articles:
                item = self.parse_product_item(article, current_url)
                if item:
                    item_url = item.get("Link chi tiết")
                    if item_url in seen_urls:
                        continue
                    seen_urls.add(item_url)
                    self.results.append(item)

            logger.info(f"-> Thu thập thành công {len(articles)} sản phẩm từ Trang {page_count}. Tổng số hiện tại: {len(self.results)}")

            # Tìm link trang kế tiếp (Next page)
            next_tag = soup.find("li", class_="next")
            if next_tag and next_tag.find("a"):
                next_rel_url = next_tag.find("a")["href"]
                current_url = urljoin(current_url, next_rel_url)
                # Polite rate-limiting giữa các trang
                delay = random.uniform(*self.delay_range)
                time.sleep(delay)
            else:
                logger.info("Không có nút 'Next'. Toàn bộ các trang đã được cào xong.")
                break

        logger.info(f"Hoàn tất quy trình cào! Tổng cộng thu thập được {len(self.results)} sản phẩm hợp lệ.")
        return self.results

    def export_csv(self, output_path: str) -> bool:
        """
        Xuất danh sách dữ liệu ra file CSV chuẩn UTF-8 có BOM (utf-8-sig).
        Giúp Microsoft Excel trên Windows/macOS mở lên hiển thị tiếng Việt hoàn hảo, không bị vỡ font.
        """
        import csv

        if not self.results:
            logger.error("Không có dữ liệu để xuất CSV.")
            return False

        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
        fieldnames = list(self.results[0].keys())

        try:
            with open(output_path, mode="w", encoding="utf-8-sig", newline="") as csvfile:
                writer = csv.DictWriter(csvfile, fieldnames=fieldnames)
                writer.writeheader()
                for row in self.results:
                    writer.writerow(row)
            logger.info(f"Đã xuất thành công file CSV: {output_path} ({len(self.results)} dòng)")
            return True
        except Exception as e:
            logger.error(f"Lỗi khi ghi file CSV {output_path}: {e}")
            return False

    def export_excel(self, output_path: str) -> bool:
        """
        Xuất danh sách dữ liệu ra bảng tính Microsoft Excel (.xlsx) với định dạng chuyên nghiệp:
        - Tiêu đề header màu xanh navy đậm, chữ trắng in đậm
        - Định dạng số tiền tệ rõ ràng
        - Tự động căn chỉnh độ rộng cột (Auto-fit Column Width)
        """
        if not HAS_OPENPYXL:
            logger.warning("Thư viện 'openpyxl' chưa được cài đặt. Bỏ qua xuất định dạng Excel nâng cao.")
            return False

        if not self.results:
            logger.error("Không có dữ liệu để xuất Excel.")
            return False

        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)

        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Báo Giá Dữ Liệu B2B"

            # 1. Định nghĩa phong cách
            header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
            header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

            data_font = Font(name="Calibri", size=10)
            center_alignment = Alignment(horizontal="center", vertical="center")
            left_alignment = Alignment(horizontal="left", vertical="center")
            right_alignment = Alignment(horizontal="right", vertical="center")

            thin_border = Border(
                left=Side(style="thin", color="D9D9D9"),
                right=Side(style="thin", color="D9D9D9"),
                top=Side(style="thin", color="D9D9D9"),
                bottom=Side(style="thin", color="D9D9D9")
            )

            # 2. Ghi Header
            fieldnames = list(self.results[0].keys())
            ws.append(fieldnames)
            ws.row_dimensions[1].height = 26

            for col_num in range(1, len(fieldnames) + 1):
                cell = ws.cell(row=1, column=col_num)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = header_alignment

            # 3. Ghi Dữ liệu
            for row_idx, item in enumerate(self.results, start=2):
                row_values = list(item.values())
                ws.append(row_values)
                ws.row_dimensions[row_idx].height = 20

                for col_idx, val in enumerate(row_values, start=1):
                    cell = ws.cell(row=row_idx, column=col_idx)
                    cell.font = data_font
                    cell.border = thin_border
                    if isinstance(val, str) and val.startswith("="):
                        cell.data_type = "s"

                    # Định dạng canh lề và số liệu theo từng cột
                    col_name = fieldnames[col_idx - 1]
                    if "Giá" in col_name:
                        cell.alignment = right_alignment
                        if "VNĐ" in col_name:
                            cell.number_format = "#,##0"
                        else:
                            cell.number_format = "#,##0.00"
                    elif "Sao" in col_name or "Mã SP" in col_name or "Tồn kho" in col_name:
                        cell.alignment = center_alignment
                    else:
                        cell.alignment = left_alignment

            # 4. Tự động căn chỉnh độ rộng cột
            for col in ws.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    val_str = str(cell.value or "")
                    if len(val_str) > max_len:
                        max_len = len(val_str)
                # Giới hạn độ rộng hợp lý
                ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 48)

            wb.save(output_path)
            logger.info(f"Đã xuất thành công file Excel (.xlsx): {output_path} ({len(self.results)} dòng)")
            return True

        except Exception as e:
            logger.error(f"Lỗi khi ghi file Excel {output_path}: {e}")
            return False


def main():
    parser = argparse.ArgumentParser(description="Hermes B2B Data Scraper & ETL Pipeline")
    parser.add_argument("--pages", type=int, default=3, help="Số trang cần cào dữ liệu (Mặc định: 3)")
    parser.add_argument("--csv", type=str, default="sample_output.csv", help="Đường dẫn file CSV đầu ra")
    parser.add_argument("--excel", type=str, default="sample_output.xlsx", help="Đường dẫn file Excel (.xlsx) đầu ra")
    parser.add_argument("--delay", type=float, default=0.5, help="Độ trễ tối thiểu giữa các requests (giây)")
    args = parser.parse_args()

    current_dir = os.path.dirname(os.path.abspath(__file__))
    csv_path = os.path.join(current_dir, args.csv) if not os.path.isabs(args.csv) else args.csv
    excel_path = os.path.join(current_dir, args.excel) if not os.path.isabs(args.excel) else args.excel

    print("=" * 70)
    print("      HERMES B2B DATA SCRAPER - PRODUCTION-READY ETL PIPELINE")
    print("=" * 70)
    print(f"Cấu hình chạy: {args.pages} trang | Delay: {args.delay}s")
    print(f"File CSV đích : {csv_path}")
    print(f"File Excel đích: {excel_path}")
    print("-" * 70)

    scraper = B2BDataScraper(delay_range=(args.delay, args.delay + 0.5))
    data = scraper.scrape_pages(max_pages=args.pages)

    if not data:
        print("[!] Không thu thập được dữ liệu nào. Vui lòng kiểm tra lại kết nối mạng.")
        sys.exit(1)

    print("-" * 70)
    # Xuất cả 2 định dạng
    csv_ok = scraper.export_csv(csv_path)
    excel_ok = scraper.export_excel(excel_path)

    if not csv_ok or not excel_ok:
        logger.error("Xuất dữ liệu thất bại.")
        sys.exit(1)

    print("=" * 70)
    print(f"[THÀNH CÔNG] Đã thu thập và chuẩn hóa hoàn tất {len(data)} bản ghi dữ liệu B2B!")
    print(f"-> CSV  : {csv_path}")
    print(f"-> Excel: {excel_path}")
    print("=" * 70)

if __name__ == "__main__":
    main()
